#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
StockCheck_msl_gen.py — генератор DATA-блоку (MSL + PH) з корпоративного xlsx.
живе доки: існує лист «АП»/«MSL» у файлі замовлення (docs/Файл для замовлення з MSL.xlsx)

wsd 12.1: єдине джерело даних — xlsx. Руками DATA-блок не правимо.
Вузол H2 (b31): 965 аптек · 83 SKU · 16 мереж.

Використання:
  python3 StockCheck_msl_gen.py <xlsx> [--out data_block.js] [--stats]

Контракт вихідних блоків (формат b30, р.1393/1479):
  MSL: {"n","c","A","B","C","D","k","p"}
  PH : {"px","area","city","addr","net","oTC","iW"}

НОРМАЛІЗАЦІЯ (H2):
  · net  — strip() + \\xa0/\\u2009/подвійні пробіли → один пробіл
           ('УАХ\\xa0' у джерелі → 'УАХ')
  · px   — рядок (ключ, не число: провідні нулі й порівняння)
  · area/city/addr — strip()
  · тири — upper().strip(), допустимі лише A/B/C/D

НЕ ЧИТАЄМО колонку «Кількіть MSL»: норма — похідна (normFor у застосунку).
Колонка використовується лише гейтом звірки (StockCheck_h2_msl_data.py).
"""
import sys, json, re, unicodedata
import openpyxl

COL = {  # індекси колонок листа «АП» (0-based)
    'px': 0, 'area': 1, 'city': 2, 'addr': 3, 'net': 4,
    'oTC': 5, 'iW': 6, 'sky': 7, 'cat': 8, 'kode': 9, 'price': 10, 'norm': 11,
}
TIERS = ('A', 'B', 'C', 'D')


def norm_txt(v):
    """Прибрати невидимі пробіли й здвоєні пробіли. NBSP → звичайний пробіл."""
    if v is None:
        return ''
    s = str(v)
    s = unicodedata.normalize('NFC', s)
    s = s.replace('\xa0', ' ').replace(' ', ' ').replace(' ', ' ')
    s = re.sub(r'\s+', ' ', s)
    return s.strip()


def norm_tier(v):
    t = norm_txt(v).upper()
    if t not in TIERS:
        raise ValueError('неприпустимий ярус: %r' % v)
    return t


def read(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    msl, seen_k = [], set()
    for r in wb['MSL'].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        k = norm_txt(r[6])
        if k in seen_k:
            raise ValueError('дубль Kode у листі MSL: %s' % k)
        seen_k.add(k)
        msl.append({
            'n': norm_txt(r[0]), 'c': norm_txt(r[1]),
            'A': int(r[2]), 'B': int(r[3]), 'C': int(r[4]), 'D': int(r[5]),
            'k': k, 'p': round(float(r[7]), 2),
        })

    ph, order = {}, []
    for r in wb['АП'].iter_rows(min_row=2, values_only=True):
        if r[COL['px']] is None:
            continue
        px = norm_txt(r[COL['px']])
        rec = {
            'px': px,
            'area': norm_txt(r[COL['area']]),
            'city': norm_txt(r[COL['city']]),
            'addr': norm_txt(r[COL['addr']]),
            'net': norm_txt(r[COL['net']]),
            'oTC': norm_tier(r[COL['oTC']]),
            'iW': norm_tier(r[COL['iW']]),
        }
        if px in ph:
            if ph[px] != rec:
                raise ValueError('той самий Proxima ID з різними атрибутами: %s' % px)
        else:
            ph[px] = rec
            order.append(px)
    return msl, [ph[p] for p in order]


def block(name, rows):
    out = ['const %s=[' % name]
    for r in rows:
        out.append('  ' + json.dumps(r, ensure_ascii=False) + ',')
    out[-1] = out[-1][:-1]          # без коми на останньому
    out.append('];')
    return '\n'.join(out)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    msl, ph = read(src)

    if '--stats' in sys.argv:
        from collections import Counter
        nets = Counter(p['net'] for p in ph)
        cities = {(p['area'], p['city']) for p in ph}
        print('SKU: %d  (OTC %d / IW %d)' % (
            len(msl), sum(1 for s in msl if s['c'] == 'OTC'),
            sum(1 for s in msl if s['c'] == 'IW')))
        print('аптек: %d · мереж: %d · пар обл+місто: %d' % (len(ph), len(nets), len(cities)))
        for n, c in nets.most_common():
            print('  %4d  %s' % (c, n))

    txt = block('MSL', msl) + '\n\n' + block('PH', ph) + '\n'
    out = None
    if '--out' in sys.argv:
        out = sys.argv[sys.argv.index('--out') + 1]
    if out:
        open(out, 'w', encoding='utf-8').write(txt)
        print('записано: %s  (%.1f КБ)' % (out, len(txt.encode('utf-8')) / 1024))
    else:
        sys.stdout.write(txt)


if __name__ == '__main__':
    main()
