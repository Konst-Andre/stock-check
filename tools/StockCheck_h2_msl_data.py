#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
StockCheck_h2_msl_data.py — гейт даних вузла H2.
живе доки: DATA-блок StockCheck генерується з корпоративного xlsx

Звіряє DATA-блок зібраного білда проти джерела на ВСІХ рядках листа «АП».
Це не smoke: перевіряється кожен рядок, не вибірка.

  python3 StockCheck_h2_msl_data.py <build.html> <src.xlsx>

Що перевіряє:
  D1  PH: кожна аптека xlsx присутня в білді, усі 7 полів збігаються
  D2  PH: у білді немає аптек, яких немає у джерелі
  D3  MSL: 83 SKU, A/B/C/D/Kode/Price збігаються
  D4  НОРМА: normFor(kode, tier своєї категорії) == «Кількіть MSL» у джерелі
      (О-16 — саме цей гейт ловить повернення однокольонкової формули)
  D5  ЦІНА: MSL.p == Custom.Price у кожному рядку
  D6  НОРМАЛІЗАЦІЯ: у полі net немає \\xa0 та здвоєних пробілів
"""
import sys, re, json, unicodedata
import openpyxl

TIERS = ('A', 'B', 'C', 'D')


def norm_txt(v):
    if v is None:
        return ''
    s = unicodedata.normalize('NFC', str(v))
    s = s.replace('\xa0', ' ').replace(' ', ' ').replace(' ', ' ')
    return re.sub(r'\s+', ' ', s).strip()


def extract(html, name):
    m = re.search(r'const %s=\[(.*?)\n\];' % name, html, re.S)
    if not m:
        raise SystemExit('✗ якір const %s=[ не знайдено' % name)
    return json.loads('[' + m.group(1).rstrip().rstrip(',') + ']')


def main():
    build, src = sys.argv[1], sys.argv[2]
    html = open(build, encoding='utf-8').read()
    MSL = extract(html, 'MSL')
    PH = extract(html, 'PH')
    by_k = {s['k']: s for s in MSL}
    by_px = {p['px']: p for p in PH}

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    fails = []

    # D3 — MSL проти джерела
    src_msl = {}
    for r in wb['MSL'].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        src_msl[norm_txt(r[6])] = r
    if len(src_msl) != len(by_k):
        fails.append('D3 кількість SKU: білд %d, джерело %d' % (len(by_k), len(src_msl)))
    for k, r in src_msl.items():
        s = by_k.get(k)
        if not s:
            fails.append('D3 SKU %s відсутній у білді' % k); continue
        if (s['A'], s['B'], s['C'], s['D']) != (r[2], r[3], r[4], r[5]):
            fails.append('D3 яруси SKU %s: білд %s, джерело %s' % (k, [s[t] for t in TIERS], r[2:6]))
        if abs(s['p'] - float(r[7])) > 0.005:
            fails.append('D3 ціна SKU %s: %s vs %s' % (k, s['p'], r[7]))

    # D1/D2/D4/D5 — прохід по всіх рядках «АП»
    seen_px, rows, bad_norm, bad_price = set(), 0, 0, 0
    for r in wb['АП'].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        rows += 1
        px = norm_txt(r[0]); seen_px.add(px)
        p = by_px.get(px)
        if not p:
            if len(fails) < 40:
                fails.append('D1 аптека %s відсутня в білді' % px)
            continue
        exp = {'area': norm_txt(r[1]), 'city': norm_txt(r[2]), 'addr': norm_txt(r[3]),
               'net': norm_txt(r[4]), 'oTC': norm_txt(r[5]).upper(), 'iW': norm_txt(r[6]).upper()}
        for f, v in exp.items():
            if p[f] != v and len(fails) < 40:
                fails.append('D1 %s.%s: білд %r, джерело %r' % (px, f, p[f], v))
        k = norm_txt(r[9]); s = by_k.get(k)
        if not s:
            continue
        tier = exp['oTC'] if norm_txt(r[8]) == 'OTC' else exp['iW']   # D4: ярус СВОЄЇ категорії
        if s[tier] != r[11]:
            bad_norm += 1
            if bad_norm <= 5:
                fails.append('D4 норма %s/%s: normFor=%s, файл=%s (кат %s, ярус %s)'
                             % (px, k, s[tier], r[11], r[8], tier))
        if abs(s['p'] - float(r[10] or 0)) > 0.005:
            bad_price += 1

    extra = set(by_px) - seen_px
    if extra:
        fails.append('D2 у білді %d аптек, яких немає у джерелі: %s' % (len(extra), sorted(extra)[:5]))
    if bad_norm:
        fails.append('D4 РАЗОМ розходжень норми: %d з %d' % (bad_norm, rows))
    if bad_price:
        fails.append('D5 розходжень ціни: %d' % bad_price)
    for p in PH:                                                       # D6
        if '\xa0' in p['net'] or '  ' in p['net'] or p['net'] != p['net'].strip():
            fails.append('D6 ненормалізований net: %r' % p['net'])

    print('рядків «АП»: %d · аптек: %d · SKU: %d' % (rows, len(by_px), len(by_k)))
    if fails:
        print('\n'.join('  ✗ ' + f for f in fails[:40]))
        print('─── ГЕЙТ ВАЛИТЬСЯ: %d зауважень ───' % len(fails))
        sys.exit(1)
    print('─── D1–D6 ✓ усі %d рядків збігаються з джерелом ───' % rows)


if __name__ == '__main__':
    main()
